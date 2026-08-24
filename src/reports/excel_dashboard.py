from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.analytics.dashboard_data import DashboardData
from src.utils.config import BASE_DIR
from src.utils.logger import logger


class ExcelDashboard:
    """
    Geração e atualização do Dashboard Executivo em Excel.

    Fluxo:

        Fato_Vendas
            ↓
        DashboardData
            ↓
        ExcelDashboard
            ↓
        Dashboard.xlsx
    """

    def __init__(self) -> None:

        self.template = (
            BASE_DIR
            / "dashboard"
            / "template"
            / "Dashboard_Template.xlsx"
        )

        self.output = (
            BASE_DIR
            / "dashboard"
            / "output"
            / "Dashboard.xlsx"
        )

    # =====================================================
    # ATUALIZAÇÃO DO DASHBOARD
    # =====================================================

    def atualizar(
        self,
        df: pd.DataFrame,
    ) -> None:
        """
        Atualiza o Dashboard utilizando os dados
        analíticos preparados pelo DashboardData.
        """

        logger.info(
            "Iniciando atualização do Dashboard Executivo."
        )

        self.output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        dados = DashboardData(df).dados()

        workbook = load_workbook(self.template)

        # -------------------------------------------------
        # BASE DE DADOS
        # -------------------------------------------------

        self._substituir_aba(
            workbook,
            "Base_Dados",
            df,
        )

        # -------------------------------------------------
        # ABAS ANALÍTICAS
        # -------------------------------------------------

        abas_dataframe = {
            "Ranking_Sellers": dados[
                "ranking_sellers"
            ],

            "Ranking_Frete": dados[
                "ranking_frete"
            ],

            "Ranking_Preco": dados[
                "ranking_preco_medio"
            ],

            "Desempenho_Estados": dados[
                "desempenho_estados"
            ],

            "Financeiro_Sellers": dados[
                "financeiro_sellers"
            ],

            "Margem_Sellers": dados[
                "ranking_margem_sellers"
            ],

            "Financeiro_Estados": dados[
                "financeiro_estados"
            ],

            "Evolucao_Mensal": dados[
                "evolucao_mensal"
            ],

            "Crescimento_Mensal": dados[
                "crescimento_mensal"
            ],

            "Ranking_Mensal": dados[
                "ranking_mensal_receita"
            ],

            "Ranking_Resultado": dados[
                "ranking_mensal_resultado"
            ],
        }

        for nome_aba, dataframe in abas_dataframe.items():

            self._substituir_aba(
                workbook,
                nome_aba,
                dataframe,
            )

        # -------------------------------------------------
        # RESUMO EXECUTIVO
        # -------------------------------------------------

        self._criar_resumo(
            workbook,
            dados,
        )

        # -------------------------------------------------
        # SALVAMENTO
        # -------------------------------------------------

        logger.info(
            "Salvando Dashboard Executivo."
        )

        workbook.save(self.output)

        logger.info(
            "Dashboard Executivo atualizado com sucesso: %s",
            self.output,
        )

    # =====================================================
    # COMPATIBILIDADE COM MÉTODO ANTIGO
    # =====================================================

    def atualizar_base(
        self,
        df: pd.DataFrame,
        aba: str = "Base_Dados",
    ) -> None:
        """
        Mantém compatibilidade com a versão anterior.
        """

        logger.info(
            "Atualizando base do Dashboard."
        )

        workbook = load_workbook(
            self.template
        )

        self._substituir_aba(
            workbook,
            aba,
            df,
        )

        self.output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook.save(
            self.output
        )

        logger.info(
            "Base do Dashboard atualizada."
        )

    # =====================================================
    # SUBSTITUIÇÃO DE ABA
    # =====================================================

    @staticmethod
    def _substituir_aba(
        workbook,
        nome_aba: str,
        df: pd.DataFrame,
    ) -> None:
        """
        Remove uma aba existente e cria uma nova
        contendo o DataFrame informado.
        """

        if nome_aba in workbook.sheetnames:

            worksheet = workbook[
                nome_aba
            ]

            workbook.remove(
                worksheet
            )

        worksheet = workbook.create_sheet(
            nome_aba
        )

        if df is None:
            return

        if isinstance(df, pd.DataFrame):

            if df.empty:
                return

            worksheet.append(
                df.columns.tolist()
            )

            for linha in df.itertuples(
                index=False,
                name=None,
            ):
                worksheet.append(
                    list(linha)
                )

    # =====================================================
    # RESUMO EXECUTIVO
    # =====================================================

    @staticmethod
    def _criar_resumo(
        workbook,
        dados: dict,
    ) -> None:
        """
        Cria a aba de resumo executivo
        com os principais KPIs.
        """

        nome_aba = "Resumo_Executivo"

        if nome_aba in workbook.sheetnames:

            workbook.remove(
                workbook[nome_aba]
            )

        worksheet = workbook.create_sheet(
            nome_aba,
            0,
        )

        indicadores = dados[
            "indicadores_financeiros"
        ]

        worksheet.append(
            [
                "INDICADOR",
                "VALOR",
            ]
        )

        linhas = [
            (
                "Receita Bruta",
                indicadores[
                    "receita_bruta"
                ],
            ),

            (
                "Frete Total",
                indicadores[
                    "frete_total"
                ],
            ),

            (
                "Resultado Após Frete",
                indicadores[
                    "resultado_apos_frete"
                ],
            ),

            (
                "Frete sobre Receita (%)",
                indicadores[
                    "frete_percentual"
                ],
            ),

            (
                "Margem Após Frete (%)",
                indicadores[
                    "margem_apos_frete"
                ],
            ),

            (
                "Total de Itens",
                indicadores[
                    "total_itens"
                ],
            ),

            (
                "Total de Pedidos",
                indicadores[
                    "total_pedidos"
                ],
            ),

            (
                "Ticket Médio por Item",
                indicadores[
                    "ticket_medio_item"
                ],
            ),

            (
                "Ticket Médio por Pedido",
                indicadores[
                    "ticket_medio_pedido"
                ],
            ),
        ]

        for linha in linhas:

            worksheet.append(
                list(linha)
            )

        # -------------------------------------------------
        # RESUMO TEMPORAL
        # -------------------------------------------------

        temporal = dados[
            "resumo_temporal"
        ]

        worksheet.append([])
        worksheet.append(
            [
                "ANÁLISE TEMPORAL",
                "VALOR",
            ]
        )

        temporal_linhas = [
            (
                "Melhor mês em receita",
                temporal.get(
                    "melhor_mes_receita"
                ),
            ),

            (
                "Receita do melhor mês",
                temporal.get(
                    "melhor_mes_receita_valor"
                ),
            ),

            (
                "Pior mês em receita",
                temporal.get(
                    "pior_mes_receita"
                ),
            ),

            (
                "Receita do pior mês",
                temporal.get(
                    "pior_mes_receita_valor"
                ),
            ),

            (
                "Melhor mês em resultado",
                temporal.get(
                    "melhor_mes_resultado"
                ),
            ),

            (
                "Resultado do melhor mês",
                temporal.get(
                    "melhor_mes_resultado_valor"
                ),
            ),

            (
                "Crescimento da receita (%)",
                temporal.get(
                    "crescimento_receita_total"
                ),
            ),

            (
                "Meses analisados",
                temporal.get(
                    "meses_analisados"
                ),
            ),
        ]

        for linha in temporal_linhas:

            worksheet.append(
                list(linha)
            )