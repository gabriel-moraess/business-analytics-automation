import os
import re
import logging

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList


class DashboardExcel:
    """
    Cria e formata o Dashboard Executivo no arquivo Excel
    gerado pelo pipeline principal.

    Compatível com o pipeline atual do
    Business Analytics Automation.

    IMPORTANTE:
    O arquivo é salvo diretamente no caminho final.
    Não são utilizados arquivos temporários.
    """

    # ============================================================
    # CORES
    # ============================================================

    AZUL = "1F4E78"
    AZUL_ESCURO = "17365D"
    AZUL_CLARO = "D9EAF7"
    AZUL_MUITO_CLARO = "F4F8FB"

    CINZA = "666666"
    CINZA_CLARO = "F2F2F2"
    CINZA_BORDA = "D9E2F3"

    BRANCO = "FFFFFF"

    VERDE = "548235"
    VERDE_CLARO = "E2F0D9"

    LARANJA = "C55A11"
    LARANJA_CLARO = "FCE4D6"

    VERMELHO = "C00000"
    VERMELHO_CLARO = "F4CCCC"

    # ============================================================
    # INICIALIZAÇÃO
    # ============================================================

    def __init__(
        self,
        caminho_excel,
        logger=None,
        *args,
        **kwargs
    ):

        self.caminho_excel = caminho_excel

        self.logger = logger or logging.getLogger(
            __name__
        )

    # ============================================================
    # CRIAR DASHBOARD
    # ============================================================

    def criar_dashboard(
        self,
        resumo=None,
        top_sellers=None,
        por_estado=None,
        fato_vendas=None,
        *args,
        **kwargs
    ):

        self.logger.info(
            "Iniciando criação do Dashboard Excel."
        )

        if not os.path.exists(
            self.caminho_excel
        ):

            raise FileNotFoundError(
                f"Arquivo Excel não encontrado: "
                f"{self.caminho_excel}"
            )

        # --------------------------------------------------------
        # Carregar workbook
        # --------------------------------------------------------

        wb = load_workbook(
            self.caminho_excel
        )

        # --------------------------------------------------------
        # Remover dashboards anteriores
        # --------------------------------------------------------

        for nome in [
            "Dashboard",
            "Dashboard_Dados"
        ]:

            if nome in wb.sheetnames:

                del wb[nome]

        # --------------------------------------------------------
        # Criar abas
        # --------------------------------------------------------

        ws = wb.create_sheet(
            "Dashboard",
            0
        )

        ws_dados = wb.create_sheet(
            "Dashboard_Dados"
        )

        # --------------------------------------------------------
        # Normalizar dados
        # --------------------------------------------------------

        resumo_df = self._normalizar_dataframe(
            resumo
        )

        top_sellers_df = self._normalizar_dataframe(
            top_sellers
        )

        estados_df = self._normalizar_dataframe(
            por_estado
        )

        fato_df = self._normalizar_dataframe(
            fato_vendas
        )

        # --------------------------------------------------------
        # KPIs
        # --------------------------------------------------------

        kpis = self._extrair_kpis(
            resumo_df,
            fato_df
        )

        # --------------------------------------------------------
        # Contexto
        # --------------------------------------------------------

        contexto = self._extrair_contexto()

        # --------------------------------------------------------
        # Cabeçalho
        # --------------------------------------------------------

        self._criar_cabecalho(
            ws,
            contexto
        )

        # --------------------------------------------------------
        # KPIs
        # --------------------------------------------------------

        self._criar_kpis(
            ws,
            kpis
        )

        # --------------------------------------------------------
        # Tabelas
        # --------------------------------------------------------

        self._criar_tabela_top_sellers(
            ws,
            top_sellers_df
        )

        self._criar_tabela_estados(
            ws,
            estados_df
        )

        # --------------------------------------------------------
        # Dados auxiliares
        # --------------------------------------------------------

        self._criar_dashboard_dados(
            ws_dados,
            resumo_df,
            top_sellers_df,
            estados_df
        )

        # --------------------------------------------------------
        # Gráficos
        # --------------------------------------------------------

        self._criar_graficos(
            ws,
            top_sellers_df,
            estados_df
        )

        # --------------------------------------------------------
        # Formatação geral
        # --------------------------------------------------------

        self._formatar_dashboard(
            ws
        )

        # --------------------------------------------------------
        # Salvar
        # --------------------------------------------------------

        self._salvar_seguro(
            wb
        )

        self.logger.info(
            "Dashboard Excel criado com sucesso."
        )

    # ============================================================
    # NORMALIZAÇÃO
    # ============================================================

    def _normalizar_dataframe(
        self,
        df
    ):

        if df is None:

            return pd.DataFrame()

        if isinstance(
            df,
            pd.Series
        ):

            return df.to_frame()

        if isinstance(
            df,
            list
        ):

            return pd.DataFrame(
                df
            )

        return df.copy()

    # ============================================================
    # CONTEXTO
    # ============================================================

    def _extrair_contexto(
        self
    ):

        nome = os.path.basename(
            self.caminho_excel
        )

        contexto = {
            "cliente": "Gabriel Tecnologia LTDA",
            "periodo": "Período automático"
        }

        # --------------------------------------------------------
        # Cliente
        # --------------------------------------------------------

        match_cliente = re.search(
            r"Business_Analytics_(.+?)_\d{4}-\d{2}_a_\d{4}-\d{2}",
            nome
        )

        if match_cliente:

            contexto[
                "cliente"
            ] = match_cliente.group(1).replace(
                "_",
                " "
            )

        # --------------------------------------------------------
        # Período
        # --------------------------------------------------------

        match_periodo = re.search(
            r"(\d{4}-\d{2})_a_(\d{4}-\d{2})",
            nome
        )

        if match_periodo:

            contexto[
                "periodo"
            ] = (
                f"{match_periodo.group(1)} "
                f"→ "
                f"{match_periodo.group(2)}"
            )

        return contexto

    # ============================================================
    # KPIs
    # ============================================================

    def _extrair_kpis(
        self,
        resumo_df,
        fato_df
    ):

        kpis = {

            "receita_bruta": 0,

            "frete_total": 0,

            "resultado_apos_frete": 0,

            "frete_percentual": 0,

            "margem_percentual": 0,

            "total_pedidos": 0,

            "total_itens": 0,

            "ticket_medio_pedido": 0,

        }

        # --------------------------------------------------------
        # Resumo
        # --------------------------------------------------------

        if not resumo_df.empty:

            linha = resumo_df.iloc[0]

            mapa = {

                "receita_bruta": [
                    "receita_bruta",
                    "Receita bruta",
                    "receita"
                ],

                "frete_total": [
                    "frete_total",
                    "Frete total",
                    "frete"
                ],

                "resultado_apos_frete": [
                    "resultado_apos_frete",
                    "Resultado após frete"
                ],

                "frete_percentual": [
                    "frete_sobre_receita",
                    "frete_percentual",
                    "Frete sobre receita"
                ],

                "margem_percentual": [
                    "margem_apos_frete",
                    "margem_percentual",
                    "Margem após frete"
                ],

                "total_pedidos": [
                    "total_pedidos",
                    "Total de pedidos",
                    "pedidos"
                ],

                "total_itens": [
                    "total_itens",
                    "Total de itens",
                    "itens"
                ],

                "ticket_medio_pedido": [
                    "ticket_medio_pedido",
                    "Ticket médio por pedido"
                ],
            }

            for destino, possibilidades in mapa.items():

                for coluna in possibilidades:

                    if coluna in linha.index:

                        valor = linha[coluna]

                        if pd.notna(valor):

                            try:

                                kpis[
                                    destino
                                ] = float(
                                    valor
                                )

                            except (
                                TypeError,
                                ValueError
                            ):

                                pass

                            break

        # --------------------------------------------------------
        # Fato_Vendas
        # --------------------------------------------------------

        if not fato_df.empty:

            colunas = set(
                fato_df.columns
            )

            # ----------------------------------------------------
            # Receita
            # ----------------------------------------------------

            if kpis[
                "receita_bruta"
            ] == 0:

                for coluna in [
                    "receita_item",
                    "price"
                ]:

                    if coluna in colunas:

                        kpis[
                            "receita_bruta"
                        ] = pd.to_numeric(
                            fato_df[coluna],
                            errors="coerce"
                        ).sum()

                        break

            # ----------------------------------------------------
            # Frete
            # ----------------------------------------------------

            if kpis[
                "frete_total"
            ] == 0:

                for coluna in [
                    "frete_item",
                    "freight_value"
                ]:

                    if coluna in colunas:

                        kpis[
                            "frete_total"
                        ] = pd.to_numeric(
                            fato_df[coluna],
                            errors="coerce"
                        ).sum()

                        break

            # ----------------------------------------------------
            # Itens
            # ----------------------------------------------------

            if kpis[
                "total_itens"
            ] == 0:

                kpis[
                    "total_itens"
                ] = len(
                    fato_df
                )

            # ----------------------------------------------------
            # Pedidos
            # ----------------------------------------------------

            if kpis[
                "total_pedidos"
            ] == 0:

                if "order_id" in colunas:

                    kpis[
                        "total_pedidos"
                    ] = fato_df[
                        "order_id"
                    ].nunique()

            # ----------------------------------------------------
            # Resultado
            # ----------------------------------------------------

            if kpis[
                "resultado_apos_frete"
            ] == 0:

                kpis[
                    "resultado_apos_frete"
                ] = (
                    kpis["receita_bruta"]
                    -
                    kpis["frete_total"]
                )

            # ----------------------------------------------------
            # Frete %
            # ----------------------------------------------------

            if kpis[
                "frete_percentual"
            ] == 0:

                if kpis[
                    "receita_bruta"
                ] != 0:

                    kpis[
                        "frete_percentual"
                    ] = (
                        kpis["frete_total"]
                        /
                        kpis["receita_bruta"]
                    )

            # ----------------------------------------------------
            # Margem
            # ----------------------------------------------------

            if kpis[
                "margem_percentual"
            ] == 0:

                if kpis[
                    "receita_bruta"
                ] != 0:

                    kpis[
                        "margem_percentual"
                    ] = (
                        kpis[
                            "resultado_apos_frete"
                        ]
                        /
                        kpis[
                            "receita_bruta"
                        ]
                    )

            # ----------------------------------------------------
            # Ticket
            # ----------------------------------------------------

            if kpis[
                "ticket_medio_pedido"
            ] == 0:

                if kpis[
                    "total_pedidos"
                ] != 0:

                    kpis[
                        "ticket_medio_pedido"
                    ] = (
                        kpis[
                            "receita_bruta"
                        ]
                        /
                        kpis[
                            "total_pedidos"
                        ]
                    )

        # --------------------------------------------------------
        # Percentuais
        # --------------------------------------------------------

        kpis[
            "frete_percentual"
        ] = self._normalizar_percentual(
            kpis[
                "frete_percentual"
            ]
        )

        kpis[
            "margem_percentual"
        ] = self._normalizar_percentual(
            kpis[
                "margem_percentual"
            ]
        )

        return kpis

    # ============================================================
    # NORMALIZAR %
    # ============================================================

    def _normalizar_percentual(
        self,
        valor
    ):

        try:

            valor = float(
                valor
            )

        except (
            TypeError,
            ValueError
        ):

            return 0

        if abs(valor) > 1:

            return valor / 100

        return valor

    # ============================================================
    # CABEÇALHO
    # ============================================================

    def _criar_cabecalho(
        self,
        ws,
        contexto
    ):

        # --------------------------------------------------------
        # Fundo
        # --------------------------------------------------------

        for row in range(
            1,
            5
        ):

            for col in range(
                1,
                16
            ):

                ws.cell(
                    row=row,
                    column=col
                ).fill = PatternFill(
                    fill_type="solid",
                    fgColor=self.AZUL_ESCURO
                )

        # --------------------------------------------------------
        # Título
        # --------------------------------------------------------

        ws.merge_cells(
            "B2:J2"
        )

        ws["B2"] = (
            "BUSINESS ANALYTICS AUTOMATION"
        )

        ws["B2"].font = Font(
            bold=True,
            size=22,
            color=self.BRANCO
        )

        ws["B2"].alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

        # --------------------------------------------------------
        # Cliente
        # --------------------------------------------------------

        ws.merge_cells(
            "K2:O2"
        )

        ws["K2"] = contexto[
            "cliente"
        ]

        ws["K2"].font = Font(
            bold=True,
            size=11,
            color=self.BRANCO
        )

        ws["K2"].alignment = Alignment(
            horizontal="right",
            vertical="center"
        )

        # --------------------------------------------------------
        # Subtítulo
        # --------------------------------------------------------

        ws.merge_cells(
            "B3:J3"
        )

        ws["B3"] = (
            "Dashboard Executivo | "
            "Desempenho Comercial, Financeiro e Operacional"
        )

        ws["B3"].font = Font(
            size=10,
            color="DDEBF7"
        )

        ws["B3"].alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

        # --------------------------------------------------------
        # Período
        # --------------------------------------------------------

        ws.merge_cells(
            "K3:O3"
        )

        ws["K3"] = (
            "PERÍODO: "
            + contexto["periodo"]
        )

        ws["K3"].font = Font(
            bold=True,
            size=10,
            color="DDEBF7"
        )

        ws["K3"].alignment = Alignment(
            horizontal="right",
            vertical="center"
        )

    # ============================================================
    # KPIs
    # ============================================================

    def _criar_kpis(
        self,
        ws,
        kpis
    ):

        cards = [

            (
                "B6",
                "RECEITA BRUTA",
                kpis[
                    "receita_bruta"
                ],
                'R$ #,##0.00'
            ),

            (
                "E6",
                "RESULTADO APÓS FRETE",
                kpis[
                    "resultado_apos_frete"
                ],
                'R$ #,##0.00'
            ),

            (
                "H6",
                "MARGEM APÓS FRETE",
                kpis[
                    "margem_percentual"
                ],
                '0.00%'
            ),

            (
                "K6",
                "TOTAL DE PEDIDOS",
                kpis[
                    "total_pedidos"
                ],
                '#,##0'
            ),

            (
                "B10",
                "FRETE TOTAL",
                kpis[
                    "frete_total"
                ],
                'R$ #,##0.00'
            ),

            (
                "E10",
                "FRETE / RECEITA",
                kpis[
                    "frete_percentual"
                ],
                '0.00%'
            ),

            (
                "H10",
                "TICKET MÉDIO",
                kpis[
                    "ticket_medio_pedido"
                ],
                'R$ #,##0.00'
            ),

            (
                "K10",
                "TOTAL DE ITENS",
                kpis[
                    "total_itens"
                ],
                '#,##0'
            ),

        ]

        for (
            celula,
            titulo,
            valor,
            formato
        ) in cards:

            coluna = ws[
                celula
            ].column

            linha = ws[
                celula
            ].row

            # ----------------------------------------------------
            # Card
            # ----------------------------------------------------

            for row in range(
                linha,
                linha + 3
            ):

                for col in range(
                    coluna,
                    coluna + 2
                ):

                    cell = ws.cell(
                        row=row,
                        column=col
                    )

                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor=self.AZUL_MUITO_CLARO
                    )

                    cell.border = Border(
                        left=Side(
                            style="thin",
                            color=self.CINZA_BORDA
                        ),
                        right=Side(
                            style="thin",
                            color=self.CINZA_BORDA
                        ),
                        top=Side(
                            style="thin",
                            color=self.CINZA_BORDA
                        ),
                        bottom=Side(
                            style="thin",
                            color=self.CINZA_BORDA
                        )
                    )

            # ----------------------------------------------------
            # Título
            # ----------------------------------------------------

            titulo_cell = ws.cell(
                row=linha,
                column=coluna
            )

            titulo_cell.value = titulo

            titulo_cell.font = Font(
                bold=True,
                size=9,
                color=self.CINZA
            )

            titulo_cell.alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

            # ----------------------------------------------------
            # Valor
            # ----------------------------------------------------

            valor_cell = ws.cell(
                row=linha + 1,
                column=coluna
            )

            valor_cell.value = valor

            valor_cell.font = Font(
                bold=True,
                size=17,
                color=self.AZUL
            )

            valor_cell.number_format = formato

            valor_cell.alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

    # ============================================================
    # TOP SELLERS
    # ============================================================

    def _criar_tabela_top_sellers(
        self,
        ws,
        df
    ):

        titulo_row = 15
        header_row = 16
        dados_row = 17

        ws.merge_cells(
            "B15:G15"
        )

        titulo = ws[
            "B15"
        ]

        titulo.value = (
            "TOP 10 SELLERS POR RECEITA"
        )

        titulo.font = Font(
            bold=True,
            size=12,
            color=self.AZUL
        )

        titulo.alignment = Alignment(
            horizontal="left"
        )

        if df.empty:

            return

        seller_col = self._encontrar_coluna(
            df,
            ["seller"]
        )

        receita_col = self._encontrar_coluna(
            df,
            ["receita"]
        )

        frete_col = self._encontrar_coluna(
            df,
            ["frete"]
        )

        margem_col = self._encontrar_coluna(
            df,
            ["margem"]
        )

        if seller_col is None:

            seller_col = df.columns[0]

        if receita_col is None:

            if len(df.columns) > 1:

                receita_col = df.columns[1]

            else:

                return

        # --------------------------------------------------------
        # Cabeçalho
        # --------------------------------------------------------

        headers = [
            "Rank",
            "Seller",
            "Receita",
            "Frete",
            "Margem"
        ]

        for col, valor in enumerate(
            headers,
            start=2
        ):

            cell = ws.cell(
                row=header_row,
                column=col
            )

            cell.value = valor

            cell.font = Font(
                bold=True,
                color=self.BRANCO,
                size=9
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.AZUL
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # --------------------------------------------------------
        # Dados
        # --------------------------------------------------------

        dados = df.copy()

        dados[
            receita_col
        ] = pd.to_numeric(
            dados[
                receita_col
            ],
            errors="coerce"
        )

        dados = dados.sort_values(
            receita_col,
            ascending=False
        ).head(10)

        for rank, (
            _,
            registro
        ) in enumerate(
            dados.iterrows(),
            start=1
        ):

            row = dados_row + rank - 1

            # Rank
            ws.cell(
                row=row,
                column=2
            ).value = rank

            # Seller
            seller = str(
                registro[
                    seller_col
                ]
            )

            ws.cell(
                row=row,
                column=3
            ).value = seller[:12]

            # Receita
            receita = pd.to_numeric(
                registro[
                    receita_col
                ],
                errors="coerce"
            )

            ws.cell(
                row=row,
                column=4
            ).value = receita

            ws.cell(
                row=row,
                column=4
            ).number_format = (
                'R$ #,##0.00'
            )

            # Frete
            if frete_col is not None:

                frete = pd.to_numeric(
                    registro[
                        frete_col
                    ],
                    errors="coerce"
                )

                ws.cell(
                    row=row,
                    column=5
                ).value = frete

                ws.cell(
                    row=row,
                    column=5
                ).number_format = (
                    'R$ #,##0.00'
                )

            # Margem
            if margem_col is not None:

                margem = pd.to_numeric(
                    registro[
                        margem_col
                    ],
                    errors="coerce"
                )

                margem = (
                    self._normalizar_percentual(
                        margem
                    )
                )

                ws.cell(
                    row=row,
                    column=6
                ).value = margem

                ws.cell(
                    row=row,
                    column=6
                ).number_format = (
                    '0.00%'
                )

            # ----------------------------------------------------
            # Linhas
            # ----------------------------------------------------

            for col in range(
                2,
                7
            ):

                cell = ws.cell(
                    row=row,
                    column=col
                )

                cell.border = Border(
                    bottom=Side(
                        style="hair",
                        color=self.CINZA_BORDA
                    )
                )

                if rank % 2 == 0:

                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F8FAFC"
                    )

    # ============================================================
    # ESTADOS
    # ============================================================

    def _criar_tabela_estados(
        self,
        ws,
        df
    ):

        titulo_row = 15
        header_row = 16
        dados_row = 17

        ws.merge_cells(
            "I15:N15"
        )

        titulo = ws[
            "I15"
        ]

        titulo.value = (
            "TOP 10 ESTADOS POR RECEITA"
        )

        titulo.font = Font(
            bold=True,
            size=12,
            color=self.AZUL
        )

        titulo.alignment = Alignment(
            horizontal="left"
        )

        if df.empty:

            return

        estado_col = self._encontrar_coluna(
            df,
            ["state", "estado"]
        )

        receita_col = self._encontrar_coluna(
            df,
            ["receita"]
        )

        if estado_col is None:

            estado_col = df.columns[0]

        if receita_col is None:

            if len(df.columns) > 1:

                receita_col = df.columns[1]

            else:

                return

        dados = df.copy()

        dados[
            receita_col
        ] = pd.to_numeric(
            dados[
                receita_col
            ],
            errors="coerce"
        )

        dados = dados.sort_values(
            receita_col,
            ascending=False
        ).head(10)

        # --------------------------------------------------------
        # Cabeçalho
        # --------------------------------------------------------

        headers = [
            "Estado",
            "Receita",
            "% Receita"
        ]

        for col, valor in enumerate(
            headers,
            start=9
        ):

            cell = ws.cell(
                row=header_row,
                column=col
            )

            cell.value = valor

            cell.font = Font(
                bold=True,
                color=self.BRANCO,
                size=9
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=self.AZUL
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        receita_total = dados[
            receita_col
        ].sum()

        # --------------------------------------------------------
        # Dados
        # --------------------------------------------------------

        for index, (
            _,
            registro
        ) in enumerate(
            dados.iterrows(),
            start=0
        ):

            row = dados_row + index

            estado = registro[
                estado_col
            ]

            receita = pd.to_numeric(
                registro[
                    receita_col
                ],
                errors="coerce"
            )

            ws.cell(
                row=row,
                column=9
            ).value = estado

            ws.cell(
                row=row,
                column=10
            ).value = receita

            ws.cell(
                row=row,
                column=10
            ).number_format = (
                'R$ #,##0.00'
            )

            if receita_total:

                percentual = (
                    receita
                    /
                    receita_total
                )

            else:

                percentual = 0

            ws.cell(
                row=row,
                column=11
            ).value = percentual

            ws.cell(
                row=row,
                column=11
            ).number_format = (
                '0.0%'
            )

            for col in range(
                9,
                12
            ):

                cell = ws.cell(
                    row=row,
                    column=col
                )

                cell.border = Border(
                    bottom=Side(
                        style="hair",
                        color=self.CINZA_BORDA
                    )
                )

                if index % 2 == 1:

                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F8FAFC"
                    )

    # ============================================================
    # ENCONTRAR COLUNA
    # ============================================================

    def _encontrar_coluna(
        self,
        df,
        termos
    ):

        for coluna in df.columns:

            nome = str(
                coluna
            ).lower()

            for termo in termos:

                if termo.lower() in nome:

                    return coluna

        return None

    # ============================================================
    # GRÁFICOS
    # ============================================================

    def _criar_graficos(
        self,
        ws,
        top_sellers_df,
        estados_df
    ):

        ws_dados = ws.parent[
            "Dashboard_Dados"
        ]

        # ========================================================
        # SELLERS
        # ========================================================

        if not top_sellers_df.empty:

            seller_col = self._encontrar_coluna(
                top_sellers_df,
                ["seller"]
            )

            receita_col = self._encontrar_coluna(
                top_sellers_df,
                ["receita"]
            )

            if (
                seller_col is not None
                and
                receita_col is not None
            ):

                dados = top_sellers_df.copy()

                dados[
                    receita_col
                ] = pd.to_numeric(
                    dados[
                        receita_col
                    ],
                    errors="coerce"
                )

                dados = dados.sort_values(
                    receita_col,
                    ascending=True
                ).head(10)

                inicio = 2

                ws_dados.cell(
                    row=inicio,
                    column=1
                ).value = "Seller"

                ws_dados.cell(
                    row=inicio,
                    column=2
                ).value = "Receita"

                for i, (
                    _,
                    registro
                ) in enumerate(
                    dados.iterrows(),
                    start=inicio + 1
                ):

                    ws_dados.cell(
                        row=i,
                        column=1
                    ).value = str(
                        registro[
                            seller_col
                        ]
                    )[:10]

                    ws_dados.cell(
                        row=i,
                        column=2
                    ).value = registro[
                        receita_col
                    ]

                chart = BarChart()

                chart.type = "bar"

                chart.style = 10

                chart.title = (
                    "Receita por Seller"
                )

                chart.height = 8.5

                chart.width = 17

                chart.legend = None

                chart.varyColors = True

                data = Reference(
                    ws_dados,
                    min_col=2,
                    min_row=inicio,
                    max_row=inicio + len(dados)
                )

                categorias = Reference(
                    ws_dados,
                    min_col=1,
                    min_row=inicio + 1,
                    max_row=inicio + len(dados)
                )

                chart.add_data(
                    data,
                    titles_from_data=True
                )

                chart.set_categories(
                    categorias
                )

                chart.dataLabels = DataLabelList()

                chart.dataLabels.showVal = True

                chart.dataLabels.showLegendKey = False

                chart.x_axis.numFmt = (
                    'R$ #,##0'
                )

                chart.gapWidth = 40

                ws.add_chart(
                    chart,
                    "B30"
                )

        # ========================================================
        # ESTADOS
        # ========================================================

        if not estados_df.empty:

            estado_col = self._encontrar_coluna(
                estados_df,
                ["state", "estado"]
            )

            receita_col = self._encontrar_coluna(
                estados_df,
                ["receita"]
            )

            if (
                estado_col is not None
                and
                receita_col is not None
            ):

                dados = estados_df.copy()

                dados[
                    receita_col
                ] = pd.to_numeric(
                    dados[
                        receita_col
                    ],
                    errors="coerce"
                )

                dados = dados.sort_values(
                    receita_col,
                    ascending=False
                ).head(10)

                inicio = 20

                ws_dados.cell(
                    row=inicio,
                    column=1
                ).value = "Estado"

                ws_dados.cell(
                    row=inicio,
                    column=2
                ).value = "Receita"

                for i, (
                    _,
                    registro
                ) in enumerate(
                    dados.iterrows(),
                    start=inicio + 1
                ):

                    ws_dados.cell(
                        row=i,
                        column=1
                    ).value = registro[
                        estado_col
                    ]

                    ws_dados.cell(
                        row=i,
                        column=2
                    ).value = registro[
                        receita_col
                    ]

                chart = BarChart()

                chart.type = "col"

                chart.style = 11

                chart.title = (
                    "Receita por Estado"
                )

                chart.height = 8.5

                chart.width = 17

                chart.legend = None

                chart.varyColors = True

                data = Reference(
                    ws_dados,
                    min_col=2,
                    min_row=inicio,
                    max_row=inicio + len(dados)
                )

                categorias = Reference(
                    ws_dados,
                    min_col=1,
                    min_row=inicio + 1,
                    max_row=inicio + len(dados)
                )

                chart.add_data(
                    data,
                    titles_from_data=True
                )

                chart.set_categories(
                    categorias
                )

                chart.dataLabels = DataLabelList()

                chart.dataLabels.showVal = True

                chart.dataLabels.showLegendKey = False

                chart.y_axis.numFmt = (
                    'R$ #,##0'
                )

                chart.gapWidth = 60

                ws.add_chart(
                    chart,
                    "I30"
                )

    # ============================================================
    # DADOS AUXILIARES
    # ============================================================

    def _criar_dashboard_dados(
        self,
        ws,
        resumo_df,
        top_sellers_df,
        estados_df
    ):

        ws.sheet_view.showGridLines = False

        linha = 1

        conjuntos = [

            (
                "Resumo",
                resumo_df
            ),

            (
                "Top_Sellers",
                top_sellers_df
            ),

            (
                "Por_Estado",
                estados_df
            ),

        ]

        for nome, df in conjuntos:

            if df.empty:

                continue

            ws.cell(
                row=linha,
                column=1
            ).value = nome

            linha += 1

            # ----------------------------------------------------
            # Cabeçalho
            # ----------------------------------------------------

            for col, coluna in enumerate(
                df.columns,
                start=1
            ):

                ws.cell(
                    row=linha,
                    column=col
                ).value = coluna

            linha += 1

            # ----------------------------------------------------
            # Dados
            # ----------------------------------------------------

            for _, registro in df.head(50).iterrows():

                for col, coluna in enumerate(
                    df.columns,
                    start=1
                ):

                    valor = registro[
                        coluna
                    ]

                    ws.cell(
                        row=linha,
                        column=col
                    ).value = (
                        None
                        if pd.isna(valor)
                        else valor
                    )

                linha += 1

            linha += 2

        ws.sheet_state = "hidden"

    # ============================================================
    # FORMATAÇÃO
    # ============================================================

    def _formatar_dashboard(
        self,
        ws
    ):

        # --------------------------------------------------------
        # Colunas
        # --------------------------------------------------------

        larguras = {

            "A": 2,

            "B": 9,
            "C": 14,
            "D": 16,

            "E": 9,
            "F": 14,
            "G": 16,

            "H": 9,
            "I": 14,
            "J": 16,

            "K": 9,
            "L": 14,
            "M": 16,

            "N": 12,
            "O": 12,
        }

        for coluna, largura in (
            larguras.items()
        ):

            ws.column_dimensions[
                coluna
            ].width = largura

        # --------------------------------------------------------
        # Altura
        # --------------------------------------------------------

        ws.row_dimensions[
            1
        ].height = 8

        ws.row_dimensions[
            2
        ].height = 30

        ws.row_dimensions[
            3
        ].height = 20

        ws.row_dimensions[
            4
        ].height = 8

        for linha in [
            6,
            7,
            8,
            10,
            11,
            12
        ]:

            ws.row_dimensions[
                linha
            ].height = 23

        ws.row_dimensions[
            15
        ].height = 25

        ws.row_dimensions[
            16
        ].height = 22

        # --------------------------------------------------------
        # Congelar
        # --------------------------------------------------------

        ws.freeze_panes = None

        # --------------------------------------------------------
        # Grid
        # --------------------------------------------------------

        ws.sheet_view.showGridLines = False

        # --------------------------------------------------------
        # Zoom
        # --------------------------------------------------------

        ws.sheet_view.zoomScale = 85

        # --------------------------------------------------------
        # Alinhamento
        # --------------------------------------------------------

        for row in ws.iter_rows():

            for cell in row:

                if cell.value is not None:

                    cell.alignment = Alignment(
                        vertical="center",
                        wrap_text=False
                    )

        # --------------------------------------------------------
        # Área de impressão
        # --------------------------------------------------------

        ws.print_area = (
            "B2:O47"
        )

        ws.page_setup.orientation = (
            "landscape"
        )

        ws.page_setup.fitToWidth = 1

        ws.page_setup.fitToHeight = 1

        ws.sheet_properties.pageSetUpPr.fitToPage = True

        # --------------------------------------------------------
        # Margens
        # --------------------------------------------------------

        ws.page_margins.left = 0.25

        ws.page_margins.right = 0.25

        ws.page_margins.top = 0.35

        ws.page_margins.bottom = 0.35

        # --------------------------------------------------------
        # Tab color
        # --------------------------------------------------------

        ws.sheet_properties.tabColor = (
            self.AZUL
        )

    # ============================================================
    # SALVAMENTO
    # ============================================================

    def _salvar_seguro(
        self,
        wb
    ):

        try:

            self.logger.info(
                "Salvando Dashboard Excel."
            )

            wb.save(
                self.caminho_excel
            )

            self.logger.info(
                f"Arquivo Excel atualizado: "
                f"{self.caminho_excel}"
            )

        except Exception as e:

            self.logger.error(
                f"Erro ao salvar o Dashboard Excel: {e}"
            )

            raise