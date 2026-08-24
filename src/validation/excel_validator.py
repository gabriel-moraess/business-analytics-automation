from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.utils.logger import logger


class ExcelValidator:
    """
    Valida automaticamente o arquivo Excel gerado pelo pipeline.

    A validação verifica:

    - existência do arquivo;
    - existência das abas esperadas;
    - presença de registros;
    - integridade das principais métricas;
    - consistência da Fato_Vendas;
    - período analisado;
    - existência do Dashboard;
    - presença de dados auxiliares.
    """

    ABAS_OBRIGATORIAS = [
        "Dashboard",
        "Resumo",
        "Top_Sellers",
        "Frete",
        "Preco_Medio",
        "Por_Estado",
        "Fato_Vendas",
    ]

    COLUNAS_FATO_OBRIGATORIAS = [
        "order_id",
        "order_item_id",
        "product_id",
        "seller_id",
        "order_purchase_timestamp",
        "price",
        "freight_value",
        "receita_item",
        "frete_item",
        "receita_total_item",
    ]

    def __init__(
        self,
        excel_path: str | Path,
        periodo_inicio: str | None = None,
        periodo_fim: str | None = None,
    ) -> None:

        self.excel_path = Path(excel_path)

        self.periodo_inicio = (
            pd.Timestamp(periodo_inicio)
            if periodo_inicio
            else None
        )

        self.periodo_fim = (
            pd.Timestamp(periodo_fim)
            if periodo_fim
            else None
        )

        self.erros: list[str] = []
        self.avisos: list[str] = []

    # =========================================================
    # MÉTODO PRINCIPAL
    # =========================================================

    def validar(self) -> bool:

        logger.info("=" * 60)
        logger.info("INICIANDO VALIDAÇÃO DO EXCEL")
        logger.info("=" * 60)

        self.erros.clear()
        self.avisos.clear()

        # -----------------------------------------------------
        # 1. ARQUIVO
        # -----------------------------------------------------

        if not self._validar_arquivo():

            self._finalizar_validacao()

            return False

        # -----------------------------------------------------
        # 2. ABAS
        # -----------------------------------------------------

        workbook = self._carregar_workbook()

        if workbook is None:

            self._finalizar_validacao()

            return False

        self._validar_abas(workbook)

        # -----------------------------------------------------
        # 3. RESUMO
        # -----------------------------------------------------

        resumo = self._carregar_aba(
            "Resumo"
        )

        if resumo is not None:

            self._validar_resumo(
                resumo
            )

        # -----------------------------------------------------
        # 4. FATO DE VENDAS
        # -----------------------------------------------------

        fato = self._carregar_aba(
            "Fato_Vendas"
        )

        if fato is not None:

            self._validar_fato_vendas(
                fato
            )

        # -----------------------------------------------------
        # 5. TOP SELLERS
        # -----------------------------------------------------

        top_sellers = self._carregar_aba(
            "Top_Sellers"
        )

        if top_sellers is not None:

            self._validar_top_sellers(
                top_sellers
            )

        # -----------------------------------------------------
        # 6. ESTADOS
        # -----------------------------------------------------

        por_estado = self._carregar_aba(
            "Por_Estado"
        )

        if por_estado is not None:

            self._validar_estados(
                por_estado
            )

        # -----------------------------------------------------
        # 7. DASHBOARD
        # -----------------------------------------------------

        self._validar_dashboard(
            workbook
        )

        # -----------------------------------------------------
        # 8. RESULTADO
        # -----------------------------------------------------

        self._finalizar_validacao()

        return len(self.erros) == 0

    # =========================================================
    # ARQUIVO
    # =========================================================

    def _validar_arquivo(self) -> bool:

        logger.info(
            "Validando existência do arquivo..."
        )

        if not self.excel_path.exists():

            self.erros.append(
                f"Arquivo não encontrado: "
                f"{self.excel_path}"
            )

            logger.error(
                "Arquivo Excel não encontrado."
            )

            return False

        if self.excel_path.stat().st_size == 0:

            self.erros.append(
                "Arquivo Excel está vazio."
            )

            logger.error(
                "Arquivo Excel está vazio."
            )

            return False

        logger.info(
            "Arquivo Excel encontrado."
        )

        logger.info(
            "Tamanho: %.2f MB",
            self.excel_path.stat().st_size
            / (1024 * 1024),
        )

        return True

    # =========================================================
    # WORKBOOK
    # =========================================================

    def _carregar_workbook(self):

        try:

            workbook = load_workbook(
                self.excel_path,
                read_only=True,
                data_only=False,
            )

            logger.info(
                "Workbook carregado para validação."
            )

            return workbook

        except Exception as erro:

            self.erros.append(
                f"Erro ao abrir workbook: {erro}"
            )

            logger.exception(
                "Erro ao abrir workbook."
            )

            return None

    # =========================================================
    # ABAS
    # =========================================================

    def _validar_abas(
        self,
        workbook,
    ) -> None:

        logger.info(
            "Validando abas obrigatórias..."
        )

        abas_existentes = workbook.sheetnames

        for aba in self.ABAS_OBRIGATORIAS:

            if aba not in abas_existentes:

                self.erros.append(
                    f"Aba obrigatória ausente: {aba}"
                )

                logger.error(
                    "Aba ausente: %s",
                    aba,
                )

            else:

                logger.info(
                    "OK - Aba encontrada: %s",
                    aba,
                )

        if "Dashboard_Dados" in abas_existentes:

            logger.info(
                "OK - Aba auxiliar Dashboard_Dados encontrada."
            )

    # =========================================================
    # CARREGAR ABA
    # =========================================================

    def _carregar_aba(
        self,
        nome_aba: str,
    ):

        try:

            dataframe = pd.read_excel(
                self.excel_path,
                sheet_name=nome_aba,
            )

            logger.info(
                "Aba '%s': %s registros.",
                nome_aba,
                len(dataframe),
            )

            return dataframe

        except Exception as erro:

            self.erros.append(
                f"Erro ao carregar aba "
                f"'{nome_aba}': {erro}"
            )

            logger.error(
                "Erro ao carregar aba '%s': %s",
                nome_aba,
                erro,
            )

            return None

    # =========================================================
    # RESUMO
    # =========================================================

    def _validar_resumo(
        self,
        resumo: pd.DataFrame,
    ) -> None:

        logger.info(
            "Validando aba Resumo..."
        )

        if resumo.empty:

            self.erros.append(
                "Aba Resumo está vazia."
            )

            return

        colunas_obrigatorias = [
            "receita_total",
            "frete_total",
            "total_itens",
            "total_sellers",
            "preco_medio_item",
        ]

        for coluna in colunas_obrigatorias:

            if coluna not in resumo.columns:

                self.erros.append(
                    f"Coluna ausente no Resumo: "
                    f"{coluna}"
                )

        if self.erros:

            return

        linha = resumo.iloc[0]

        # -----------------------------------------------------
        # RECEITA
        # -----------------------------------------------------

        receita = linha["receita_total"]

        if pd.isna(receita):

            self.erros.append(
                "Receita total está nula."
            )

        elif receita < 0:

            self.erros.append(
                "Receita total possui valor negativo."
            )

        else:

            logger.info(
                "OK - Receita total: R$ %.2f",
                receita,
            )

        # -----------------------------------------------------
        # FRETE
        # -----------------------------------------------------

        frete = linha["frete_total"]

        if pd.isna(frete):

            self.erros.append(
                "Frete total está nulo."
            )

        elif frete < 0:

            self.erros.append(
                "Frete total possui valor negativo."
            )

        else:

            logger.info(
                "OK - Frete total: R$ %.2f",
                frete,
            )

        # -----------------------------------------------------
        # ITENS
        # -----------------------------------------------------

        total_itens = linha["total_itens"]

        if pd.isna(total_itens):

            self.erros.append(
                "Total de itens está nulo."
            )

        elif total_itens <= 0:

            self.erros.append(
                "Total de itens deve ser maior que zero."
            )

        else:

            logger.info(
                "OK - Total de itens: %s",
                total_itens,
            )

        # -----------------------------------------------------
        # SELLERS
        # -----------------------------------------------------

        total_sellers = linha["total_sellers"]

        if pd.isna(total_sellers):

            self.erros.append(
                "Total de sellers está nulo."
            )

        elif total_sellers <= 0:

            self.erros.append(
                "Total de sellers deve ser maior que zero."
            )

        else:

            logger.info(
                "OK - Sellers ativos: %s",
                total_sellers,
            )

        # -----------------------------------------------------
        # PREÇO MÉDIO
        # -----------------------------------------------------

        preco_medio = linha[
            "preco_medio_item"
        ]

        if pd.isna(preco_medio):

            self.erros.append(
                "Preço médio está nulo."
            )

        elif preco_medio < 0:

            self.erros.append(
                "Preço médio possui valor negativo."
            )

        else:

            logger.info(
                "OK - Preço médio: R$ %.2f",
                preco_medio,
            )

    # =========================================================
    # FATO DE VENDAS
    # =========================================================

    def _validar_fato_vendas(
        self,
        fato: pd.DataFrame,
    ) -> None:

        logger.info(
            "Validando Fato_Vendas..."
        )

        # -----------------------------------------------------
        # VAZIO
        # -----------------------------------------------------

        if fato.empty:

            self.erros.append(
                "Fato_Vendas está vazia."
            )

            return

        logger.info(
            "OK - Fato_Vendas possui %s registros.",
            len(fato),
        )

        # -----------------------------------------------------
        # COLUNAS
        # -----------------------------------------------------

        for coluna in self.COLUNAS_FATO_OBRIGATORIAS:

            if coluna not in fato.columns:

                self.erros.append(
                    f"Coluna obrigatória ausente "
                    f"em Fato_Vendas: {coluna}"
                )

        if self.erros:

            return

        # -----------------------------------------------------
        # IDS NULOS
        # -----------------------------------------------------

        colunas_ids = [
            "order_id",
            "product_id",
            "seller_id",
        ]

        for coluna in colunas_ids:

            nulos = fato[coluna].isna().sum()

            if nulos > 0:

                self.erros.append(
                    f"Existem {nulos} valores nulos "
                    f"em '{coluna}'."
                )

            else:

                logger.info(
                    "OK - Sem nulos em %s.",
                    coluna,
                )

        # -----------------------------------------------------
        # RECEITA
        # -----------------------------------------------------

        receita_nula = fato[
            "receita_item"
        ].isna().sum()

        if receita_nula > 0:

            self.erros.append(
                f"Existem {receita_nula} valores "
                "nulos em receita_item."
            )

        else:

            logger.info(
                "OK - receita_item sem nulos."
            )

        # -----------------------------------------------------
        # FRETE
        # -----------------------------------------------------

        frete_nulo = fato[
            "frete_item"
        ].isna().sum()

        if frete_nulo > 0:

            self.erros.append(
                f"Existem {frete_nulo} valores "
                "nulos em frete_item."
            )

        else:

            logger.info(
                "OK - frete_item sem nulos."
            )

        # -----------------------------------------------------
        # VALORES NEGATIVOS
        # -----------------------------------------------------

        receita_negativa = (
            fato["receita_item"] < 0
        ).sum()

        frete_negativo = (
            fato["frete_item"] < 0
        ).sum()

        if receita_negativa > 0:

            self.erros.append(
                f"{receita_negativa} registros "
                "possuem receita negativa."
            )

        if frete_negativo > 0:

            self.erros.append(
                f"{frete_negativo} registros "
                "possuem frete negativo."
            )

        if (
            receita_negativa == 0
            and frete_negativo == 0
        ):

            logger.info(
                "OK - Não existem receitas "
                "ou fretes negativos."
            )

        # -----------------------------------------------------
        # DATAS
        # -----------------------------------------------------

        fato[
            "order_purchase_timestamp"
        ] = pd.to_datetime(
            fato[
                "order_purchase_timestamp"
            ],
            errors="coerce",
        )

        datas_invalidas = fato[
            "order_purchase_timestamp"
        ].isna().sum()

        if datas_invalidas > 0:

            self.erros.append(
                f"{datas_invalidas} datas "
                "de pedido inválidas."
            )

        else:

            logger.info(
                "OK - Datas de pedido válidas."
            )

        # -----------------------------------------------------
        # PERÍODO
        # -----------------------------------------------------

        self._validar_periodo(
            fato
        )

        # -----------------------------------------------------
        # DUPLICIDADES
        # -----------------------------------------------------

        duplicados = fato.duplicated(
            subset=[
                "order_id",
                "order_item_id",
            ]
        ).sum()

        if duplicados > 0:

            self.erros.append(
                f"{duplicados} registros duplicados "
                "em order_id + order_item_id."
            )

        else:

            logger.info(
                "OK - Não existem duplicidades "
                "em order_id + order_item_id."
            )

    # =========================================================
    # PERÍODO
    # =========================================================

    def _validar_periodo(
        self,
        fato: pd.DataFrame,
    ) -> None:

        if (
            self.periodo_inicio is None
            or self.periodo_fim is None
        ):

            logger.info(
                "Período esperado não informado; "
                "validação de período ignorada."
            )

            return

        datas = fato[
            "order_purchase_timestamp"
        ]

        menor_data = datas.min()
        maior_data = datas.max()

        logger.info(
            "Período encontrado: %s até %s",
            menor_data.strftime("%Y-%m-%d"),
            maior_data.strftime("%Y-%m-%d"),
        )

        if menor_data < self.periodo_inicio:

            self.erros.append(
                "Existem registros anteriores "
                "ao período solicitado."
            )

        if maior_data > self.periodo_fim:

            self.erros.append(
                "Existem registros posteriores "
                "ao período solicitado."
            )

        if (
            menor_data >= self.periodo_inicio
            and maior_data <= self.periodo_fim
        ):

            logger.info(
                "OK - Registros dentro do período esperado."
            )

    # =========================================================
    # TOP SELLERS
    # =========================================================

    def _validar_top_sellers(
        self,
        dataframe: pd.DataFrame,
    ) -> None:

        logger.info(
            "Validando Top_Sellers..."
        )

        if dataframe.empty:

            self.erros.append(
                "Top_Sellers está vazia."
            )

            return

        if len(dataframe) > 10:

            self.avisos.append(
                "Top_Sellers possui mais de 10 registros."
            )

        if "receita" not in dataframe.columns:

            self.erros.append(
                "Coluna 'receita' ausente "
                "em Top_Sellers."
            )

            return

        if (
            dataframe["receita"]
            .isna()
            .any()
        ):

            self.erros.append(
                "Top_Sellers possui receitas nulas."
            )

        if (
            dataframe["receita"] < 0
        ).any():

            self.erros.append(
                "Top_Sellers possui receitas negativas."
            )

        logger.info(
            "OK - Top_Sellers validado."
        )

    # =========================================================
    # ESTADOS
    # =========================================================

    def _validar_estados(
        self,
        dataframe: pd.DataFrame,
    ) -> None:

        logger.info(
            "Validando Por_Estado..."
        )

        if dataframe.empty:

            self.erros.append(
                "Por_Estado está vazia."
            )

            return

        if "receita" not in dataframe.columns:

            self.erros.append(
                "Coluna 'receita' ausente "
                "em Por_Estado."
            )

            return

        if (
            dataframe["receita"]
            .isna()
            .any()
        ):

            self.erros.append(
                "Por_Estado possui receitas nulas."
            )

        if (
            dataframe["receita"] < 0
        ).any():

            self.erros.append(
                "Por_Estado possui receitas negativas."
            )

        logger.info(
            "OK - Por_Estado validado."
        )

    # =========================================================
    # DASHBOARD
    # =========================================================

    def _validar_dashboard(
        self,
        workbook,
    ) -> None:

        logger.info(
            "Validando Dashboard..."
        )

        if "Dashboard" not in workbook.sheetnames:

            self.erros.append(
                "Dashboard não encontrado."
            )

            return

        dashboard = workbook[
            "Dashboard"
        ]

        # -----------------------------------------------------
        # TÍTULO
        # -----------------------------------------------------

        titulo = dashboard["B2"].value

        if not titulo:

            self.erros.append(
                "Dashboard não possui título."
            )

        else:

            logger.info(
                "OK - Dashboard possui título."
            )

        # -----------------------------------------------------
        # VISIBILIDADE
        # -----------------------------------------------------

        if dashboard.sheet_state != "visible":

            self.erros.append(
                "Dashboard não está visível."
            )

        else:

            logger.info(
                "OK - Dashboard está visível."
            )

        # -----------------------------------------------------
        # GRÁFICOS
        # -----------------------------------------------------

        quantidade_graficos = len(
            dashboard._charts
        )

        logger.info(
            "Gráficos encontrados: %s",
            quantidade_graficos,
        )

        if quantidade_graficos < 4:

            self.avisos.append(
                "Dashboard possui menos de "
                "4 gráficos."
            )

        else:

            logger.info(
                "OK - Dashboard possui "
                "%s gráficos.",
                quantidade_graficos,
            )

    # =========================================================
    # FINALIZAÇÃO
    # =========================================================

    def _finalizar_validacao(
        self,
    ) -> None:

        logger.info("=" * 60)
        logger.info("RESULTADO DA VALIDAÇÃO")
        logger.info("=" * 60)

        # -----------------------------------------------------
        # ERROS
        # -----------------------------------------------------

        if self.erros:

            logger.error(
                "VALIDAÇÃO REPROVADA"
            )

            logger.error(
                "Total de erros: %s",
                len(self.erros),
            )

            for erro in self.erros:

                logger.error(
                    "ERRO: %s",
                    erro,
                )

        else:

            logger.info(
                "VALIDAÇÃO APROVADA"
            )

            logger.info(
                "Nenhum erro crítico encontrado."
            )

        # -----------------------------------------------------
        # AVISOS
        # -----------------------------------------------------

        if self.avisos:

            logger.warning(
                "Total de avisos: %s",
                len(self.avisos),
            )

            for aviso in self.avisos:

                logger.warning(
                    "AVISO: %s",
                    aviso,
                )

        logger.info("=" * 60)